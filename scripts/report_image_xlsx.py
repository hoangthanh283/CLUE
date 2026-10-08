"""Build the image-CL (ViT) report workbook: per-scenario comparison sheets with charts.

Mirrors the per-combo tabs of the Google "Reports" sheet (AA bar, BWT bar, line chart of
mean seen-task accuracy after each task). Import the xlsx into Sheets via
File > Import > Insert new sheet(s) — charts survive the import.

    uv run --with openpyxl python scripts/report_image_xlsx.py [--out image_report.xlsx]
"""

import argparse
import glob
import json
import os
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

SCENARIOS = ["cil_cifar100", "cil_imagenet_r"]
MAX_LINE_SERIES = 8


def load_runs() -> pd.DataFrame:
    rows = []
    for sc in SCENARIOS:
        for d in sorted(glob.glob(f"results/{sc}_*")):
            m = os.path.join(d, "metrics.json")
            if not os.path.exists(m):
                continue
            j = json.load(open(m))
            name, me, sd = os.path.basename(d), j["method"], j["seed"]
            var = name.replace(f"{sc}_", "", 1).replace(f"_seed{sd}", "").replace("_vit", "")
            var = "" if var == me else re.sub(f"^{me}_?", "", var)
            M = np.array(j["matrix"], dtype=float)
            curve = [float(np.nanmean(M[t, : t + 1])) for t in range(M.shape[0])]
            rows.append(
                dict(
                    scenario=sc,
                    method_label=f"{me} [{var}]" if var else me,
                    seed=sd,
                    AA=j["AA"],
                    BWT=j["BWT"],
                    AF=j["AF"],
                    curve=curve,
                )
            )
    return pd.DataFrame(rows)


def write_scenario_sheet(wb: Workbook, df: pd.DataFrame, sc: str) -> None:
    ws = wb.create_sheet(f"{sc} (vit)")
    sub = df[df.scenario == sc]
    g = sub.groupby("method_label")
    summ = pd.DataFrame(
        {
            "n": g.size(),
            "AA": g.AA.mean(),
            "AA_std": g.AA.std(),
            "BWT": g.BWT.mean(),
            "AF": g.AF.mean(),
        }
    ).reset_index()
    summ = summ.sort_values("AA", ascending=False).round(2)
    ws.append([f"{sc} on vit — CL method comparison (mean over seeds)"])
    ws.append([])
    ws.append(["method_label", "n", "AA", "AA_std", "BWT", "AF"])
    for r in summ.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
    first, last = 4, 3 + len(summ)

    # Per-task mean seen-task accuracy (mean over seeds), top-k by AA + joint/naive anchors.
    cur = {
        ml: np.nanmean(np.stack(grp.curve.tolist()), axis=0)
        for ml, grp in sub.groupby("method_label")
    }
    order = [m for m in summ.method_label if m in cur]
    anchors = [m for m in ("joint", "naive") if m in order]
    series = anchors + [m for m in order if m not in anchors][: MAX_LINE_SERIES - len(anchors)]
    T = len(next(iter(cur.values())))
    crow = last + 3
    ws.cell(
        row=crow - 1, column=1, value="Mean seen-task accuracy after each task (mean over seeds)"
    )
    ws.append([])  # spacer so the curve header lands on crow
    ws.cell(row=crow, column=1, value="Task")
    for j, m in enumerate(series, start=2):
        ws.cell(row=crow, column=j, value=m)
    for t in range(T):
        ws.cell(row=crow + 1 + t, column=1, value=t + 1)
        for j, m in enumerate(series, start=2):
            v = cur[m][t]
            ws.cell(row=crow + 1 + t, column=j, value=None if np.isnan(v) else round(float(v), 2))

    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    for col, title, anchor in (
        (3, "AA (final average accuracy)", "I3"),
        (5, "BWT (backward transfer)", "I23"),
    ):
        ch = BarChart()
        ch.type, ch.title, ch.y_axis.title = "col", f"{sc} — {title}", title.split(" ")[0]
        ch.add_data(
            Reference(ws, min_col=col, min_row=first - 1, max_row=last), titles_from_data=True
        )
        ch.set_categories(cats)
        ch.width, ch.height, ch.legend = 28, 9, None
        ws.add_chart(ch, anchor)
    lc = LineChart()
    lc.title, lc.y_axis.title, lc.x_axis.title = (
        f"{sc} — mean seen-task accuracy after each task",
        "accuracy",
        "task",
    )
    lc.add_data(
        Reference(ws, min_col=2, max_col=1 + len(series), min_row=crow, max_row=crow + T),
        titles_from_data=True,
    )
    lc.set_categories(Reference(ws, min_col=1, min_row=crow + 1, max_row=crow + T))
    lc.width, lc.height = 28, 11
    ws.add_chart(lc, "I43")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="results/image_report.xlsx")
    args = ap.parse_args()
    df = load_runs()
    wb = Workbook()
    wb.remove(wb.active)
    for sc in SCENARIOS:
        write_scenario_sheet(wb, df, sc)
    wb.save(args.out)
    print(f"{len(df)} runs -> {args.out}")


if __name__ == "__main__":
    main()
